"""Word (.docx) export: a branded, professionally typeset legal document.

The document text is parsed once with :func:`parse_document` and every block is
mapped onto native Word constructs rather than pasted in as raw text:
restyled built-in styles (so Word's navigation pane and table of contents
work), real bullets, hanging-indent clauses, PAGE/NUMPAGES fields in the
footer and a repeating-header table for the terms section.
"""

from __future__ import annotations

import io
import re
from datetime import datetime, timezone

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.text.font import CT_RPr
from docx.oxml.text.parfmt import CT_PPr
from docx.section import Section, _Footer
from docx.shared import Emu, Inches, Length, Mm, Pt, RGBColor
from docx.styles.style import ParagraphStyle
from docx.table import Table, _Cell, _Row
from docx.text.paragraph import Paragraph

from backend.services.branding import (
    BRAND_NAME,
    FOOTER_DISCLAIMER,
    GOLD,
    INK,
    MUTED,
    NAVY,
    RULE,
    TABLE_HEADER_FILL,
    TABLE_STRIPE_FILL,
    Logo,
    load_default_logo,
)
from backend.services.document_service import (
    Block,
    BlockType,
    is_terms_section,
    parse_document,
    split_lead,
    split_placeholders,
)

RGB = tuple[int, int, int]

FONT_NAME = "Times New Roman"
WHITE: RGB = (255, 255, 255)

_PAGE_SIZES = {"A4": (Mm(210), Mm(297)), "LETTER": (Inches(8.5), Inches(11))}
_MARGIN = Inches(1)
_HEADER_FOOTER_DISTANCE = Inches(0.5)
_LOGO_MAX_WIDTH = Inches(2.3)
_LOGO_MAX_HEIGHT = Inches(0.8)

_INDENT_STEP = Inches(0.5)       # clause text position / nesting step
_BULLET_HANG = Inches(0.25)      # bullet glyph sits this far left of the text
_MAX_CLAUSE_LEVEL = 3
_SIGNATURE_WIDTH = Inches(2.8)
_TERM_NUMBER_WIDTH = Inches(0.55)
_BULLET_STYLES = ("List Bullet", "List Bullet 2", "List Bullet 3")
_LIST_TYPES = (BlockType.BULLET, BlockType.NUMBERED)
_MAX_PROPERTY_LENGTH = 255       # limit python-docx enforces on core properties

# Children that follow a given element in the OOXML schema, so new elements can
# be inserted in a position Word accepts.
_PBDR_SUCCESSORS = (
    "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
    "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd",
    "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
    "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
    "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
)
_RPR_SPACING_SUCCESSORS = (
    "w:w", "w:kern", "w:position", "w:sz", "w:szCs", "w:highlight", "w:u", "w:effect", "w:bdr",
    "w:shd", "w:fitText", "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang", "w:eastAsianLayout",
    "w:specVanish", "w:oMath",
)
_TBL_BORDERS_SUCCESSORS = ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription")
_TBL_CELL_MAR_SUCCESSORS = ("w:tblLook", "w:tblCaption", "w:tblDescription")
_TC_SHD_SUCCESSORS = ("w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark")

# Characters XML 1.0 cannot represent (lone surrogates, U+FFFE/U+FFFF, C0 controls).
_XML_ILLEGAL_RE = re.compile("[^\t\n\r\u0020-\ud7ff\ue000-\ufffd\U00010000-\U0010ffff]")


# ---------------------------------------------------------------------------
# Low-level OOXML helpers
# ---------------------------------------------------------------------------


def _rgb(color: RGB) -> RGBColor:
    return RGBColor(*color)


def _hex(color: RGB) -> str:
    return "{:02X}{:02X}{:02X}".format(*color)


def _set_font(rpr: CT_RPr, name: str = FONT_NAME) -> None:
    """Pin every script slot to ``name`` and drop theme fonts, which would take precedence."""
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(attr), None)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)


def _set_tracking(rpr: CT_RPr, twentieths_of_point: int) -> None:
    """Set letter spacing (character tracking)."""
    for existing in rpr.findall(qn("w:spacing")):
        rpr.remove(existing)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:val"), str(twentieths_of_point))
    rpr.insert_element_before(spacing, *_RPR_SPACING_SUCCESSORS)


def _border(edge: str, color: RGB, *, size: int, space: int = 0) -> OxmlElement:
    """A single-line border element; ``size`` is in eighths of a point, ``space`` in points."""
    element = OxmlElement(f"w:{edge}")
    element.set(qn("w:val"), "single")
    element.set(qn("w:sz"), str(size))
    element.set(qn("w:space"), str(space))
    element.set(qn("w:color"), _hex(color))
    return element


def _set_paragraph_border(ppr: CT_PPr, edge: str, color: RGB, *, size: int, space: int) -> None:
    for existing in ppr.findall(qn("w:pBdr")):
        ppr.remove(existing)
    borders = OxmlElement("w:pBdr")
    borders.append(_border(edge, color, size=size, space=space))
    ppr.insert_element_before(borders, *_PBDR_SUCCESSORS)


def _add_field_char(paragraph: Paragraph, char_type: str) -> None:
    element = OxmlElement("w:fldChar")
    element.set(qn("w:fldCharType"), char_type)
    paragraph.add_run()._r.append(element)


def _add_field(paragraph: Paragraph, instruction: str) -> None:
    """Append a field (e.g. ``PAGE``) that Word recalculates whenever it lays out pages."""
    _add_field_char(paragraph, "begin")
    code = OxmlElement("w:instrText")
    code.set(qn("xml:space"), "preserve")
    code.text = f" {instruction} "
    paragraph.add_run()._r.append(code)
    _add_field_char(paragraph, "separate")
    paragraph.add_run("1")  # cached result, shown until Word updates the field
    _add_field_char(paragraph, "end")


def _add_inline(paragraph: Paragraph, text: str, *, allow_lead: bool = True, size: Length | None = None) -> None:
    """Add text runs: a leading legal phrase in bold and [PLACEHOLDERS] highlighted."""
    lead, rest = split_lead(text) if allow_lead else ("", text)
    runs = []
    if lead:
        runs.append(paragraph.add_run(lead))
        runs[-1].bold = True
    for segment, is_placeholder in split_placeholders(rest):
        runs.append(paragraph.add_run(segment))
        if is_placeholder:
            runs[-1].font.highlight_color = WD_COLOR_INDEX.YELLOW
    if size is not None:
        for run in runs:
            run.font.size = size


# ---------------------------------------------------------------------------
# Document set-up: page, styles, package, header and footer
# ---------------------------------------------------------------------------


def _configure_page(section: Section, page_size: str) -> None:
    section.page_width, section.page_height = _PAGE_SIZES.get(str(page_size).upper(), _PAGE_SIZES["A4"])
    section.top_margin = section.bottom_margin = _MARGIN
    section.left_margin = section.right_margin = _MARGIN
    section.header_distance = section.footer_distance = _HEADER_FOOTER_DISTANCE


def _style_font(style: ParagraphStyle, *, size: float, color: RGB, bold: bool | None = None) -> None:
    _set_font(style.element.get_or_add_rPr())
    style.font.size = Pt(size)
    style.font.color.rgb = _rgb(color)
    if bold is not None:
        style.font.bold = bold


def _configure_styles(document: DocxDocument) -> None:
    """Restyle the built-in styles the renderer uses into the LegalEase house style."""
    styles = document.styles
    # Theme fonts (Calibri/Cambria) in the defaults and in any style a user applies later.
    for rfonts in list(styles.element.iter(qn("w:rFonts"))):
        if any(attr.lower().endswith("theme") for attr in rfonts.attrib):
            _set_font(rfonts.getparent())

    normal = styles["Normal"]
    _style_font(normal, size=12, color=INK)
    fmt = normal.paragraph_format
    fmt.space_before, fmt.space_after = Pt(0), Pt(6)
    fmt.line_spacing = 1.15
    fmt.widow_control = True

    title = styles["Title"]
    _style_font(title, size=16, color=NAVY, bold=True)
    title.font.all_caps = True
    _set_tracking(title.element.get_or_add_rPr(), 20)
    fmt = title.paragraph_format
    fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt.space_before, fmt.space_after = Pt(0), Pt(20)
    fmt.line_spacing = 1.1
    fmt.keep_with_next = True
    _set_paragraph_border(title.element.get_or_add_pPr(), "bottom", GOLD, size=8, space=8)

    heading = styles["Heading 1"]
    _style_font(heading, size=12, color=NAVY, bold=True)
    fmt = heading.paragraph_format
    fmt.space_before, fmt.space_after = Pt(14), Pt(6)
    fmt.line_spacing = 1.15
    fmt.keep_with_next = fmt.keep_together = True

    subheading = styles["Heading 2"]
    _style_font(subheading, size=12, color=INK, bold=True)
    fmt = subheading.paragraph_format
    fmt.space_before, fmt.space_after = Pt(10), Pt(4)
    fmt.line_spacing = 1.15
    fmt.keep_with_next = fmt.keep_together = True

    for name in _BULLET_STYLES:
        bullet = styles[name]
        _style_font(bullet, size=12, color=INK)
        # Spacing between items follows the source text instead of being suppressed.
        ppr = bullet.element.get_or_add_pPr()
        for element in ppr.findall(qn("w:contextualSpacing")):
            ppr.remove(element)

    for name in ("Header", "Footer"):
        style = styles[name]
        _style_font(style, size=9, color=MUTED)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.line_spacing = 1.0


def _configure_package(document: DocxDocument) -> None:
    """Open in Word's current layout mode (not "Compatibility Mode"); drop the template's blank thumbnail."""
    for setting in document.settings.element.xpath("./w:compat/w:compatSetting[@w:name='compatibilityMode']"):
        setting.set(qn("w:val"), "15")
    rels = document.part.package.rels
    for r_id in [r_id for r_id, rel in rels.items() if rel.reltype == RT.THUMBNAIL]:
        del rels[r_id]


def _add_logo(paragraph: Paragraph, logo: Logo | None) -> None:
    """Centre the logo (scaled to fit the header box) or, without one, a text wordmark."""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(12)  # keeps the title clear of the logo
    if logo is None:
        wordmark = paragraph.add_run(BRAND_NAME)
        wordmark.bold = True
        wordmark.font.size = Pt(22)
        wordmark.font.color.rgb = _rgb(NAVY)
        return
    width = min(_LOGO_MAX_WIDTH, round(_LOGO_MAX_HEIGHT * logo.aspect_ratio))
    height = round(width / logo.aspect_ratio)
    picture = paragraph.add_run().add_picture(logo.stream(), width=Emu(width), height=Emu(height))
    picture._inline.docPr.set("descr", "Logo")  # alt text for screen readers


def _fill_footer(footer: _Footer, *, include_disclaimer: bool) -> None:
    """Rule, optional disclaimer and a live "Page X of Y" line."""
    first = footer.paragraphs[0]
    _set_paragraph_border(first._p.get_or_add_pPr(), "top", RULE, size=4, space=6)
    if include_disclaimer:
        first.alignment = WD_ALIGN_PARAGRAPH.CENTER
        note = first.add_run(FOOTER_DISCLAIMER)
        note.italic = True
        note.font.size = Pt(8)
        page_line = footer.add_paragraph(style="Footer")
        page_line.paragraph_format.space_before = Pt(3)
    else:
        page_line = first
    page_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    page_line.add_run("Page ")
    _add_field(page_line, "PAGE")
    page_line.add_run(" of ")
    _add_field(page_line, "NUMPAGES")


def _build_header_footer(section: Section, title: str, logo: Logo | None, *, include_disclaimer: bool) -> None:
    """Logo header on page 1, running title header on later pages, footer on every page."""
    section.different_first_page_header_footer = True
    _add_logo(section.first_page_header.paragraphs[0], logo)

    running = section.header.paragraphs[0]
    running.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    running.add_run(title)
    _set_paragraph_border(running._p.get_or_add_pPr(), "bottom", RULE, size=4, space=4)

    for footer in (section.first_page_footer, section.footer):
        _fill_footer(footer, include_disclaimer=include_disclaimer)


def _set_core_properties(document: DocxDocument, title: str, doc_type: str) -> None:
    props = document.core_properties
    now = datetime.now(timezone.utc)
    props.title = title[:_MAX_PROPERTY_LENGTH]
    props.subject = doc_type[:_MAX_PROPERTY_LENGTH]
    props.author = props.last_modified_by = BRAND_NAME
    props.comments = FOOTER_DISCLAIMER
    props.created = props.modified = now
    props.revision = 1


# ---------------------------------------------------------------------------
# Block renderers
# ---------------------------------------------------------------------------


def _text_width(document: DocxDocument) -> int:
    section = document.sections[-1]
    return section.page_width - section.left_margin - section.right_margin


def _add_title(document: DocxDocument, title: str) -> None:
    _add_inline(document.add_paragraph(style="Title"), title, allow_lead=False)


def _add_heading(document: DocxDocument, block: Block) -> Paragraph:
    paragraph = document.add_paragraph(style="Heading 1")
    _add_inline(paragraph, block.display_text, allow_lead=False)
    return paragraph


def _add_subheading(document: DocxDocument, block: Block) -> Paragraph:
    paragraph = document.add_paragraph(style="Heading 2")
    _add_inline(paragraph, block.display_text, allow_lead=False)
    return paragraph


def _add_paragraph(document: DocxDocument, block: Block) -> Paragraph:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _add_inline(paragraph, block.text)
    return paragraph


def _add_bullet(document: DocxDocument, block: Block) -> Paragraph:
    level = min(block.level, len(_BULLET_STYLES) - 1)
    paragraph = document.add_paragraph(style=_BULLET_STYLES[level])
    fmt = paragraph.paragraph_format
    fmt.left_indent = Emu(_INDENT_STEP * (level + 1))
    fmt.first_line_indent = Emu(-_BULLET_HANG)
    _add_inline(paragraph, block.text)
    return paragraph


def _add_numbered(document: DocxDocument, block: Block) -> Paragraph:
    """Hanging-indent clause: the marker is authoritative, so Word numbering is not used."""
    text_start = Emu(_INDENT_STEP * (min(block.level, _MAX_CLAUSE_LEVEL) + 1))
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt = paragraph.paragraph_format
    fmt.left_indent = text_start
    fmt.first_line_indent = Emu(-_INDENT_STEP)
    fmt.tab_stops.add_tab_stop(text_start)
    marker = paragraph.add_run(f"{block.marker}\t")
    marker.bold = True
    marker.font.color.rgb = _rgb(NAVY)
    _add_inline(paragraph, block.text)
    return paragraph


def _add_signature_line(document: DocxDocument, block: Block) -> Paragraph:
    paragraph = document.add_paragraph()
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(24)  # room to sign
    fmt.right_indent = Emu(max(_text_width(document) - _SIGNATURE_WIDTH, 0))
    fmt.keep_with_next = True
    _set_paragraph_border(paragraph._p.get_or_add_pPr(), "bottom", INK, size=6, space=1)
    return paragraph


_RENDERERS = {
    BlockType.HEADING: _add_heading,
    BlockType.SUBHEADING: _add_subheading,
    BlockType.PARAGRAPH: _add_paragraph,
    BlockType.BULLET: _add_bullet,
    BlockType.NUMBERED: _add_numbered,
    BlockType.SIGNATURE_LINE: _add_signature_line,
}


# ---------------------------------------------------------------------------
# Terms table
# ---------------------------------------------------------------------------


def _is_term_item(block: Block) -> bool:
    return block.type in _LIST_TYPES and block.level == 0 and is_terms_section(block.section)


def _format_table(table: Table, width: int) -> None:
    """Full-width fixed layout, light rules and comfortable cell padding."""
    tbl_pr = table._tbl.tblPr
    table_width = tbl_pr.find(qn("w:tblW"))
    table_width.set(qn("w:type"), "dxa")
    table_width.set(qn("w:w"), str(Emu(width).twips))

    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(_border(edge, RULE, size=4))
    tbl_pr.insert_element_before(borders, *_TBL_BORDERS_SUCCESSORS)

    margins = OxmlElement("w:tblCellMar")
    for edge, twips in (("top", 72), ("left", 115), ("bottom", 72), ("right", 115)):
        margin = OxmlElement(f"w:{edge}")
        margin.set(qn("w:w"), str(twips))
        margin.set(qn("w:type"), "dxa")
        margins.append(margin)
    tbl_pr.insert_element_before(margins, *_TBL_CELL_MAR_SUCCESSORS)

    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False


def _shade(cell: _Cell, color: RGB) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), _hex(color))
    cell._tc.get_or_add_tcPr().insert_element_before(shading, *_TC_SHD_SUCCESSORS)


def _add_row(
    table: Table, *, header: bool = False, fill: RGB | None = None, keep_with_next: bool = False
) -> tuple[Paragraph, Paragraph]:
    """Add a row that never splits across pages; returns the paragraph of each cell."""
    row: _Row = table.add_row()
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))
    if header:
        tr_pr.append(OxmlElement("w:tblHeader"))  # repeat at the top of every page
    paragraphs = []
    for cell, alignment in zip(row.cells, (WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT)):
        if fill is not None:
            _shade(cell, fill)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = alignment
        fmt = paragraph.paragraph_format
        fmt.space_before = fmt.space_after = Pt(0)
        fmt.line_spacing = 1.1
        fmt.keep_with_next = keep_with_next
        paragraphs.append(paragraph)
    return paragraphs[0], paragraphs[1]


def _add_terms_table(document: DocxDocument, items: list[Block], numbered_before: int) -> None:
    """Render a run of term items as a numbered two-column table.

    ``numbered_before`` counts the terms of the same section already shown in an
    earlier table, so numbering continues (and the header is not repeated) when
    nested items split a section's terms into several tables.
    """
    text_width = _text_width(document)
    table = document.add_table(rows=0, cols=2)
    _format_table(table, text_width)
    for column, width in zip(table.columns, (_TERM_NUMBER_WIDTH, text_width - _TERM_NUMBER_WIDTH)):
        column.width = Emu(width)

    if numbered_before == 0:
        header_cells = _add_row(table, header=True, fill=TABLE_HEADER_FILL, keep_with_next=True)
        for paragraph, label in zip(header_cells, ("No.", "Term / Condition")):
            run = paragraph.add_run(label)
            run.bold = True
            run.font.size = Pt(11)
            run.font.color.rgb = _rgb(WHITE)

    last = len(items) - 1
    for offset, block in enumerate(items):
        position = numbered_before + offset
        number_cell, term_cell = _add_row(
            table,
            fill=TABLE_STRIPE_FILL if position % 2 else None,
            # Like widow/orphan control: never leave a lone first or last row on a page.
            keep_with_next=offset < last and offset in (0, last - 1),
        )
        number = number_cell.add_run(block.marker or str(position + 1))
        number.bold = True
        number.font.size = Pt(11)
        number.font.color.rgb = _rgb(NAVY)
        _add_inline(term_cell, block.text, size=Pt(11))


# ---------------------------------------------------------------------------
# Body layout
# ---------------------------------------------------------------------------


def _comparable(text: str) -> str:
    return re.sub(r"[^0-9a-z]+", " ", text.casefold()).strip()


def _without_repeated_title(blocks: list[Block], title: str) -> list[Block]:
    """Drop a first line that merely repeats the title, so the title appears once."""
    key = _comparable(title)
    if blocks and key and _comparable(blocks[0].display_text) == key:
        return blocks[1:]
    return blocks


def _keeps_with_next(block: Block, following: Block) -> bool:
    """Keep lead-ins, signature blocks and short tight groups on the page of what follows."""
    if following.type is BlockType.SIGNATURE_LINE or block.text.endswith(":"):
        return True
    return block.type is BlockType.PARAGRAPH and not following.space_before and len(block.text) <= 80


def _render_blocks(document: DocxDocument, blocks: list[Block]) -> None:
    """Append the body blocks, grouping terms into tables and spacing blocks as in the source."""
    previous_block: Block | None = None
    previous_paragraph: Paragraph | None = None  # None right after a table
    terms_shown: dict[str, int] = {}
    index = 0
    while index < len(blocks):
        block = blocks[index]

        if _is_term_item(block):
            end = index
            while end < len(blocks) and _is_term_item(blocks[end]) and blocks[end].section == block.section:
                end += 1
            shown = terms_shown.get(block.section, 0)
            _add_terms_table(document, blocks[index:end], shown)
            terms_shown[block.section] = shown + end - index
            previous_block, previous_paragraph = blocks[end - 1], None
            index = end
            continue

        paragraph = _RENDERERS[block.type](document, block)
        fmt = paragraph.paragraph_format
        after_table = previous_block is not None and previous_paragraph is None
        if after_table and block.type is not BlockType.HEADING:
            fmt.space_before = Pt(8)  # breathing room below a table
        elif previous_paragraph is not None and not block.space_before and BlockType.HEADING not in (
            block.type, previous_block.type
        ):
            # Lines written directly under each other in the source stay tight.
            list_pair = block.type in _LIST_TYPES and previous_block.type in _LIST_TYPES
            previous_paragraph.paragraph_format.space_after = Pt(2 if list_pair else 0)
            if block.type is BlockType.SUBHEADING:
                fmt.space_before = Pt(0)

        following = blocks[index + 1] if index + 1 < len(blocks) else None
        if following is not None and _keeps_with_next(block, following):
            fmt.keep_with_next = True

        previous_block, previous_paragraph = block, paragraph
        index += 1


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------


def format_docx(
    text: str,
    doc_type: str = "Legal Document",
    *,
    include_disclaimer: bool = True,
    logo: Logo | None = None,
    page_size: str = "A4",
) -> bytes:
    """Render the (edited) document text as a branded, typeset Word document.

    Args:
        text: Document text in LegalEase's plain-text conventions (Markdown is tolerated).
        doc_type: Document type, used as the fallback title and the file's subject.
        include_disclaimer: Print the AI-draft disclaimer in the footer of every page.
        logo: Header logo for page 1; ``None`` uses the LegalEase logo, or a text
            wordmark when that asset is missing.
        page_size: ``"A4"`` or ``"LETTER"`` (anything else falls back to A4).

    Returns:
        The .docx file contents.
    """
    doc_type = _XML_ILLEGAL_RE.sub("", doc_type or "").strip()
    parsed = parse_document(_XML_ILLEGAL_RE.sub("", text or ""), doc_type or "Legal Document")

    document = Document()
    section = document.sections[0]
    _configure_page(section, page_size)
    _configure_styles(document)
    _configure_package(document)
    _build_header_footer(
        section, parsed.title, logo or load_default_logo(), include_disclaimer=include_disclaimer
    )

    _add_title(document, parsed.title)
    _render_blocks(document, _without_repeated_title(parsed.blocks, parsed.title))
    _set_core_properties(document, parsed.title, doc_type or parsed.title)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
