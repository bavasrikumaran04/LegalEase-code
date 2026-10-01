"""Tests for the TXT, DOCX and PDF exporters and the /export API.

Files are re-opened with python-docx and pypdf so every assertion is made on
what a user would actually download, not on renderer internals.
"""

from __future__ import annotations

import base64
import io
import os
import re

import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Mm
from fastapi.testclient import TestClient
from PIL import Image
from pypdf import PdfReader

from backend.services import export_service, pdf_generator
from backend.services.branding import BRAND_NAME, DISCLAIMER_TEXT, FOOTER_DISCLAIMER, decode_logo
from backend.services.docx_generator import format_docx
from backend.services.pdf_generator import format_pdf
from backend.services.txt_generator import format_txt

DOC_TYPE = "Service Agreement"
TITLE = "MASTER SERVICES AGREEMENT"
EDITED_SENTENCE = "The Parties confirm that the zebra-striped quokka clause was added during editing."

SAMPLE_DOCUMENT = f"""\
{TITLE}

This Master Services Agreement (the “Agreement”) is made on April 15, 2025 between Asha Rao (the “Provider”) and Orbit Labs Pvt. Ltd. (the “Client”).

RECITALS

WHEREAS, the Client wishes to engage the Provider for software services at [CLIENT ADDRESS]; and

NOW, THEREFORE, the Parties agree as follows:

1. DEFINITIONS

1.1 “Services” means the services described in Schedule A.

1.2 “Fee” means ₹1,50,000 payable under Section 4.

2. SCOPE OF SERVICES

2.1 The Provider shall perform the Services subject to the following:
(a) the Services shall meet the agreed specifications;
(b) the Provider shall report progress weekly; and
(c) the Client shall provide timely access to [SYSTEMS].

3. TERMS AND CONDITIONS

• The Provider shall deliver the Services by May 15, 2025.
• The Client shall pay each invoice within thirty (30) days.
• Either Party may terminate on fifteen (15) days’ written notice.
• Confidentiality obligations survive termination.

4. FEES AND PAYMENT

4.1 The Client shall pay the Fee of ₹1,50,000 in two instalments. {EDITED_SENTENCE}

SIGNATURES

IN WITNESS WHEREOF, the Parties have executed this Agreement as of the date above.

____________________________
Asha Rao (Provider)

____________________________
Orbit Labs Pvt. Ltd. (Client)
Title: [TITLE]
"""

MARKDOWN_DOCUMENT = """\
# Non-Disclosure Agreement

## 1. Purpose

**Confidential Information** means *any* information disclosed — see https://example.com/"""\
    + "a" * 200 + """.

- First item
  - Nested item
- Second item costs €600

| Term | Value |
|------|-------|
| Duration | 2 years |
"""


def flat(text: str) -> str:
    """Collapse whitespace so wrapped or justified output can be searched."""
    return " ".join(text.split())


def png_logo(size: tuple[int, int] = (240, 240)) -> bytes:
    image = Image.new("RGBA", size, (200, 30, 30, 255))
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def long_document(sections: int = 30) -> str:
    parts = ["LONG SERVICE AGREEMENT", ""]
    for n in range(1, sections + 1):
        parts += [f"{n}. SECTION NUMBER {n}", ""]
        for k in range(1, 4):
            parts += [f"{n}.{k} " + "The Parties agree to perform their obligations diligently. " * 6, ""]
    return "\n".join(parts)


def docx_text(data: bytes) -> str:
    document = Document(io.BytesIO(data))
    texts = [p.text for p in document.paragraphs]
    texts += [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join(texts)


def pdf_pages(data: bytes) -> list[str]:
    return [page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages]


def footer_xml(data: bytes) -> str:
    section = Document(io.BytesIO(data)).sections[0]
    return section.first_page_footer._element.xml + section.footer._element.xml


# --------------------------------------------------------------------------- TXT


def test_txt_returns_title_first_and_edited_content() -> None:
    output = format_txt(SAMPLE_DOCUMENT, DOC_TYPE)
    assert isinstance(output, str)
    assert output.lstrip().startswith(TITLE)
    assert EDITED_SENTENCE in flat(output)
    assert "₹1,50,000" in output


def test_txt_disclaimer_toggle() -> None:
    marker = flat(DISCLAIMER_TEXT)[:40]
    assert marker in flat(format_txt(SAMPLE_DOCUMENT, DOC_TYPE, include_disclaimer=True))
    assert marker not in flat(format_txt(SAMPLE_DOCUMENT, DOC_TYPE, include_disclaimer=False))


# -------------------------------------------------------------------------- DOCX


@pytest.fixture(scope="module")
def docx_bytes() -> bytes:
    return format_docx(SAMPLE_DOCUMENT, DOC_TYPE)


def test_docx_is_a_zip_package_python_docx_can_open(docx_bytes: bytes) -> None:
    assert isinstance(docx_bytes, bytes)
    assert docx_bytes[:2] == b"PK"
    Document(io.BytesIO(docx_bytes))


def test_docx_title_appears_once_and_edit_is_present(docx_bytes: bytes) -> None:
    document = Document(io.BytesIO(docx_bytes))
    paragraphs = [p.text.strip() for p in document.paragraphs]
    assert sum(1 for text in paragraphs if text.upper() == TITLE) == 1
    assert EDITED_SENTENCE in flat(docx_text(docx_bytes))
    assert '(the "Agreement")' in docx_text(docx_bytes)
    assert "₹1,50,000" in docx_text(docx_bytes)


def test_docx_terms_section_is_a_table(docx_bytes: bytes) -> None:
    tables = Document(io.BytesIO(docx_bytes)).tables
    assert len(tables) == 1
    rows = [[cell.text.strip() for cell in row.cells] for row in tables[0].rows]
    assert rows[0] == ["No.", "Term / Condition"]
    assert len(rows) == 1 + 4
    assert [row[0] for row in rows[1:]] == ["1", "2", "3", "4"]
    assert "Confidentiality obligations survive termination." in rows[4][1]


def test_docx_bullets_use_list_styles() -> None:
    data = format_docx("NOTICE\n\nThe following apply:\n\n• First point here\n• Second point here\n", DOC_TYPE)
    bullets = [p for p in Document(io.BytesIO(data)).paragraphs if "point here" in p.text]
    assert len(bullets) == 2
    assert all(p.style.name.startswith("List Bullet") for p in bullets)


def test_docx_normal_font_is_times_new_roman(docx_bytes: bytes) -> None:
    assert Document(io.BytesIO(docx_bytes)).styles["Normal"].font.name == "Times New Roman"


def test_docx_first_page_header_has_logo(docx_bytes: bytes) -> None:
    section = Document(io.BytesIO(docx_bytes)).sections[0]
    assert section.different_first_page_header_footer
    header = section.first_page_header
    assert header._element.xpath(".//pic:pic")
    assert any(rel.reltype == RT.IMAGE for rel in header.part.rels.values())


def test_docx_footer_has_page_fields(docx_bytes: bytes) -> None:
    xml = footer_xml(docx_bytes)
    assert re.search(r"<w:instrText[^>]*>\s*PAGE\b", xml)
    assert re.search(r"<w:instrText[^>]*>\s*NUMPAGES\b", xml)


def test_docx_disclaimer_toggle() -> None:
    with_disclaimer = footer_xml(format_docx(SAMPLE_DOCUMENT, DOC_TYPE, include_disclaimer=True))
    without = footer_xml(format_docx(SAMPLE_DOCUMENT, DOC_TYPE, include_disclaimer=False))
    assert FOOTER_DISCLAIMER in with_disclaimer
    assert FOOTER_DISCLAIMER not in without


@pytest.mark.parametrize(("page_size", "width_mm", "height_mm"), [("A4", 210, 297), ("LETTER", 215.9, 279.4)])
def test_docx_page_sizes(page_size: str, width_mm: float, height_mm: float) -> None:
    section = Document(io.BytesIO(format_docx(SAMPLE_DOCUMENT, DOC_TYPE, page_size=page_size))).sections[0]
    assert abs(section.page_width - Mm(width_mm)) < Mm(1)
    assert abs(section.page_height - Mm(height_mm)) < Mm(1)


def test_docx_accepts_custom_logo() -> None:
    data = format_docx(SAMPLE_DOCUMENT, DOC_TYPE, logo=decode_logo(png_logo()))
    header = Document(io.BytesIO(data)).sections[0].first_page_header
    assert header._element.xpath(".//pic:pic")


def test_docx_core_properties(docx_bytes: bytes) -> None:
    properties = Document(io.BytesIO(docx_bytes)).core_properties
    assert properties.title == TITLE
    assert properties.author == BRAND_NAME


@pytest.mark.parametrize("text", [TITLE, MARKDOWN_DOCUMENT], ids=["title-only", "markdown"])
def test_docx_handles_edge_inputs(text: str) -> None:
    data = format_docx(text, DOC_TYPE)
    Document(io.BytesIO(data))


# --------------------------------------------------------------------------- PDF


@pytest.fixture(scope="module")
def pdf_bytes() -> bytes:
    return format_pdf(SAMPLE_DOCUMENT, DOC_TYPE)


def test_pdf_is_readable_and_contains_content(pdf_bytes: bytes) -> None:
    assert isinstance(pdf_bytes, bytes)
    assert pdf_bytes[:4] == b"%PDF"
    text = flat("\n".join(pdf_pages(pdf_bytes)))
    assert EDITED_SENTENCE in text
    assert TITLE in text
    assert "Page 1 of" in text


def test_pdf_long_document_has_page_footer_on_every_page() -> None:
    pages = pdf_pages(format_pdf(long_document(), DOC_TYPE))
    total = len(pages)
    assert total > 1
    for number, text in enumerate(pages, start=1):
        assert f"Page {number} of {total}" in flat(text)


def test_pdf_disclaimer_toggle() -> None:
    marker = flat(FOOTER_DISCLAIMER)[:30]
    assert marker in flat(pdf_pages(format_pdf(SAMPLE_DOCUMENT, DOC_TYPE))[0])
    assert marker not in flat("".join(pdf_pages(format_pdf(SAMPLE_DOCUMENT, DOC_TYPE, include_disclaimer=False))))


def test_pdf_signature_underscores_are_intact() -> None:
    text = SAMPLE_DOCUMENT + "\nSignature: ____________________   Date: ____________\n"
    joined = flat("\n".join(pdf_pages(format_pdf(text, DOC_TYPE))))
    assert "Signature: ____________________ Date: ____________" in joined
    # Stand-alone signature lines are drawn as rules; their captions must follow.
    assert "Asha Rao (Provider)" in joined
    assert "Orbit Labs Pvt. Ltd. (Client) Title: [TITLE]" in joined


def test_pdf_metadata(pdf_bytes: bytes) -> None:
    metadata = PdfReader(io.BytesIO(pdf_bytes)).metadata
    assert metadata.title == TITLE
    assert metadata.author == BRAND_NAME


@pytest.mark.parametrize(("page_size", "width", "height"), [("A4", 595.28, 841.89), ("LETTER", 612, 792)])
def test_pdf_page_sizes(page_size: str, width: float, height: float) -> None:
    box = PdfReader(io.BytesIO(format_pdf(SAMPLE_DOCUMENT, DOC_TYPE, page_size=page_size))).pages[0].mediabox
    assert abs(float(box.width) - width) < 1
    assert abs(float(box.height) - height) < 1


def test_pdf_accepts_custom_logo() -> None:
    data = format_pdf(SAMPLE_DOCUMENT, DOC_TYPE, logo=decode_logo(png_logo((300, 120))))
    assert list(PdfReader(io.BytesIO(data)).pages[0].images)


def test_pdf_core_font_fallback_handles_non_latin1(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(pdf_generator, "_find_serif_fonts", lambda: None)
    text = SAMPLE_DOCUMENT + "\nExtra: €600, → arrow, 中文 and \U0001f600.\n"
    data = format_pdf(text, DOC_TYPE)
    joined = flat("\n".join(pdf_pages(data)))
    assert EDITED_SENTENCE in joined
    assert "Rs.1,50,000" in joined


def test_pdf_handles_markdown_input() -> None:
    assert format_pdf(MARKDOWN_DOCUMENT, DOC_TYPE)[:4] == b"%PDF"


# --------------------------------------------------------------------------- API


def export_body(**overrides: object) -> dict:
    return {"content": SAMPLE_DOCUMENT, "document_type": DOC_TYPE, **overrides}


@pytest.mark.parametrize("fmt", ["docx", "pdf"])
def test_api_export_returns_file_with_edits(client: TestClient, fmt: str) -> None:
    response = client.post(f"/export/{fmt}", json=export_body())
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == export_service.MEDIA_TYPES[fmt]
    assert response.headers["content-disposition"] == f'attachment; filename="service_agreement.{fmt}"'
    text = docx_text(response.content) if fmt == "docx" else "\n".join(pdf_pages(response.content))
    assert EDITED_SENTENCE in flat(text)


@pytest.mark.parametrize("fmt", ["docx", "pdf"])
def test_api_export_accepts_base64_logo(client: TestClient, fmt: str) -> None:
    logo = base64.b64encode(png_logo()).decode()
    response = client.post(f"/export/{fmt}", json=export_body(logo_base64=logo, include_disclaimer=False))
    assert response.status_code == 200, response.text


@pytest.mark.parametrize("fmt", ["docx", "pdf"])
def test_api_export_rejects_non_image_logo(client: TestClient, fmt: str) -> None:
    logo = base64.b64encode(os.urandom(512)).decode()
    response = client.post(f"/export/{fmt}", json=export_body(logo_base64=logo))
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "invalid_logo"


def test_api_export_failure_is_generic_500(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    def explode(*args: object, **kwargs: object) -> bytes:
        raise RuntimeError("secret internal detail")

    monkeypatch.setattr(export_service, "format_pdf", explode)
    response = client.post("/export/pdf", json=export_body())
    assert response.status_code == 500
    error = response.json()["error"]
    assert error["code"] == "pdf_export_failed"
    assert error["message"] == "The PDF file could not be generated. Please try again."
    assert "secret internal detail" not in response.text
    assert "Traceback" not in response.text


# -------------------------------------------------------------- Edit round trip


def test_generated_document_edits_survive_every_export(client: TestClient, fake_generator: object) -> None:
    generated = client.post(
        "/generate",
        json={
            "document_type": "Freelance Work Contract",
            "parties": "Jane Doe (Service Provider), TechNova Inc. (Client)",
            "terms": "Payment within 30 days; Termination with 15 days notice",
            "dates": "2025-04-15",
        },
    )
    assert generated.status_code == 200, generated.text
    original = generated.json()["document"]
    assert "Jane Doe" in original

    appended = "The Parties further agree that all notices shall be sent by registered post."
    edited = original.replace("Jane Doe", "Priya Sharma").rstrip() + "\n\n" + appended + "\n"
    body = {"content": edited, "document_type": "Freelance Work Contract"}

    outputs = {}
    for fmt in ("txt", "docx", "pdf"):
        response = client.post(f"/export/{fmt}", json=body)
        assert response.status_code == 200, response.text
        outputs[fmt] = response.content
    texts = {
        "txt": outputs["txt"].decode("utf-8"),
        "docx": docx_text(outputs["docx"]),
        "pdf": "\n".join(pdf_pages(outputs["pdf"])),
    }
    for fmt, text in texts.items():
        joined = flat(text)
        assert appended in joined, fmt
        assert "Priya Sharma" in joined, fmt
        assert "Jane Doe" not in joined, fmt
